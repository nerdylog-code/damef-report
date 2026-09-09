from __future__ import annotations

from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .parser import DamefRecord

HEADERS = [
    "Empresa", "Identificador", "Inscricao_Estadual", "Ano_Base", "Municipio", "UF",
    "Tipo", "CNAE", "Entradas", "Vendas_Estado", "Vendas_Outros_Estados",
    "Vendas_Exterior", "Saidas", "VAF", "Arquivo",
]


def _number(value: Decimal) -> float:
    return float(value)


def build_workbook(records: list[DamefRecord], output_path: Path) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Consolidado"
    sheet.append(["DAMEF - relatorio consolidado"])
    sheet.append(HEADERS)
    for item in records:
        sheet.append([
            item.company, item.document_id, item.state_registration, item.base_year,
            item.municipality, item.state, item.taxpayer_type, item.economic_activity,
            _number(item.total_inputs), _number(item.domestic_sales),
            _number(item.interstate_sales), _number(item.export_sales),
            _number(item.total_outputs), _number(item.value_added), item.source_file,
        ])

    summary = workbook.create_sheet("Resumo")
    summary.append(["Resumo por empresa e ano"])
    summary.append(["Empresa", "Ano_Base", "Documentos", "Total_Saidas", "Total_Entradas", "VAF"])
    grouped: dict[tuple[str, int], list[DamefRecord]] = {}
    for item in records:
        grouped.setdefault((item.company, item.base_year), []).append(item)
    for (company, year), items in sorted(grouped.items()):
        summary.append([
            company, year, len(items),
            sum((_number(item.total_outputs) for item in items), 0.0),
            sum((_number(item.total_inputs) for item in items), 0.0),
            sum((_number(item.value_added) for item in items), 0.0),
        ])

    for current in workbook.worksheets:
        for cell in current[1]:
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.font = Font(color="FFFFFF", bold=True)
        for cell in current[2]:
            cell.fill = PatternFill("solid", fgColor="D9EAF7")
            cell.font = Font(bold=True)
        current.freeze_panes = "A3"
        current.auto_filter.ref = current.dimensions
        for column in current.columns:
            width = max(len(str(cell.value or "")) for cell in column) + 2
            current.column_dimensions[get_column_letter(column[0].column)].width = min(width, 32)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return output_path
