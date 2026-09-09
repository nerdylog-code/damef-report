from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from damef_report.parser import parse_damef_text
from damef_report.workbook import build_workbook


def test_workbook_has_consolidated_and_summary(tmp_path: Path) -> None:
    record = parse_damef_text("Contribuinte: TESTE\nAno Base: 2024\nTOTAL DAS SAÍDAS 10,00")
    output = build_workbook([record], tmp_path / "report.xlsx")
    workbook = load_workbook(output, data_only=True)
    assert workbook.sheetnames == ["Consolidado", "Resumo"]
    assert workbook["Consolidado"].cell(3, 13).value == 10
    assert workbook["Resumo"].cell(3, 4).value == 10
    workbook.close()
